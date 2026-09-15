"""Export reconciliation report to CSV or XLSX."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side as XlSide
from openpyxl.utils import get_column_letter

from .engine import apply_match_status, build_report
from .models import Match, MatchStatus, PeriodMeta, Side, Transaction


def _money(cents: int | None) -> float | None:
    if cents is None:
        return None
    return round(cents / 100.0, 2)


def _tx_rows(transactions: list[Transaction]) -> list[dict[str, Any]]:
    rows = []
    for t in sorted(transactions, key=lambda x: (x.date, x.side.value, x.description)):
        rows.append(
            {
                "side": t.side.value,
                "date": t.date,
                "description": t.description,
                "amount": _money(t.amount_cents),
                "reference": t.reference or "",
                "status": t.status.value if hasattr(t.status, "value") else t.status,
                "match_id": t.match_id or "",
                "source_file": t.source_file or "",
            }
        )
    return rows


def _match_rows(matches: list[Match], by_id: dict[str, Transaction]) -> list[dict[str, Any]]:
    rows = []
    for m in matches:
        bank_desc = " | ".join(
            by_id[i].description for i in m.bank_ids if i in by_id
        )
        book_desc = " | ".join(
            by_id[i].description for i in m.book_ids if i in by_id
        )
        rows.append(
            {
                "match_id": m.id,
                "status": m.status.value if hasattr(m.status, "value") else m.status,
                "method": m.method.value if hasattr(m.method, "value") else m.method,
                "confidence": m.confidence,
                "bank_total": _money(m.bank_total_cents),
                "book_total": _money(m.book_total_cents),
                "bank_lines": len(m.bank_ids),
                "book_lines": len(m.book_ids),
                "bank_descriptions": bank_desc,
                "book_descriptions": book_desc,
                "reasons": "; ".join(m.reasons or []),
            }
        )
    return rows


def _summary_rows(report: dict, period: PeriodMeta) -> list[tuple[str, Any]]:
    return [
        ("Company", period.company or ""),
        ("Account", period.account_label or ""),
        ("Period start", period.start_date or ""),
        ("Period end", period.end_date or ""),
        ("Exported at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Bank lines", report.get("bank_count", 0)),
        ("Book lines", report.get("book_count", 0)),
        ("Matched groups", report.get("matched_count", 0)),
        ("Proposed groups", report.get("proposed_count", 0)),
        ("Unmatched bank lines", report.get("unmatched_bank_count", 0)),
        ("Unmatched book lines", report.get("unmatched_book_count", 0)),
        ("", ""),
        ("Bank activity total", _money(report.get("bank_activity_cents"))),
        ("Book activity total", _money(report.get("book_activity_cents"))),
        ("Activity difference (book − bank)", _money(report.get("activity_difference_cents"))),
        ("", ""),
        ("Deposits in transit (unmatched book credits)", _money(report.get("deposits_in_transit_cents"))),
        ("Outstanding payments (unmatched book debits)", _money(report.get("outstanding_payments_cents"))),
        ("Unrecorded bank credits", _money(report.get("unrecorded_bank_credits_cents"))),
        ("Unrecorded bank charges", _money(report.get("unrecorded_bank_charges_cents"))),
        ("Unmatched bank sum", _money(report.get("unmatched_bank_cents"))),
        ("Unmatched book sum", _money(report.get("unmatched_book_cents"))),
        ("Explained difference", _money(report.get("explained_difference_cents"))),
        ("Difference fully explained", "Yes" if report.get("ties") else "No"),
    ]


def build_export_payload(
    transactions: list[Transaction],
    matches: list[Match],
    period: PeriodMeta,
) -> dict[str, Any]:
    txs = apply_match_status(transactions, matches)
    report = build_report(txs, matches)
    by_id = {t.id: t for t in txs}
    unmatched_bank = [t for t in txs if t.side == Side.BANK and t.status != MatchStatus.MATCHED]
    unmatched_book = [t for t in txs if t.side == Side.BOOK and t.status != MatchStatus.MATCHED]
    return {
        "summary": _summary_rows(report, period),
        "transactions": _tx_rows(txs),
        "matches": _match_rows(matches, by_id),
        "unmatched_bank": _tx_rows(unmatched_bank),
        "unmatched_book": _tx_rows(unmatched_book),
    }


def export_csv_bytes(
    transactions: list[Transaction],
    matches: list[Match],
    period: PeriodMeta,
) -> bytes:
    data = build_export_payload(transactions, matches, period)
    buf = io.StringIO()
    w = csv.writer(buf)

    w.writerow(["Concord — Bank Reconciliation Report"])
    w.writerow([])
    w.writerow(["Summary"])
    w.writerow(["Item", "Value"])
    for label, value in data["summary"]:
        w.writerow([label, value if value is not None else ""])

    def write_section(title: str, rows: list[dict[str, Any]]) -> None:
        w.writerow([])
        w.writerow([title])
        if not rows:
            w.writerow(["(none)"])
            return
        headers = list(rows[0].keys())
        w.writerow(headers)
        for row in rows:
            w.writerow([row.get(h, "") for h in headers])

    write_section("All transactions", data["transactions"])
    write_section("Matches", data["matches"])
    write_section("Unmatched bank", data["unmatched_bank"])
    write_section("Unmatched book", data["unmatched_book"])

    return buf.getvalue().encode("utf-8-sig")


def export_xlsx_bytes(
    transactions: list[Transaction],
    matches: list[Match],
    period: PeriodMeta,
) -> bytes:
    data = build_export_payload(transactions, matches, period)
    wb = Workbook()

    header_font = Font(name="Calibri", bold=True, size=12, color="1A1714")
    title_font = Font(name="Calibri", bold=True, size=14, color="1E3D34")
    thin = Border(
        left=XlSide(style="thin", color="D4CFC5"),
        right=XlSide(style="thin", color="D4CFC5"),
        top=XlSide(style="thin", color="D4CFC5"),
        bottom=XlSide(style="thin", color="D4CFC5"),
    )
    header_fill = PatternFill("solid", fgColor="E6E0D6")
    money_formats = {"amount", "bank_total", "book_total"}

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Concord — Bank Reconciliation Report"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")
    ws["A3"] = "Item"
    ws["B3"] = "Value"
    for col in ("A", "B"):
        ws[f"{col}3"].font = header_font
        ws[f"{col}3"].fill = header_fill
        ws[f"{col}3"].border = thin
    for i, (label, value) in enumerate(data["summary"], start=4):
        ws.cell(i, 1, label).border = thin
        cell = ws.cell(i, 2, value if value is not None else "")
        cell.border = thin
        if isinstance(value, float):
            cell.number_format = '#,##0.00'
        cell.alignment = Alignment(horizontal="right" if isinstance(value, (int, float)) else "left")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 22

    def add_sheet(name: str, rows: list[dict[str, Any]]) -> None:
        sheet = wb.create_sheet(name)
        if not rows:
            sheet["A1"] = "(none)"
            return
        headers = list(rows[0].keys())
        for c, h in enumerate(headers, start=1):
            cell = sheet.cell(1, c, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin
        for r, row in enumerate(rows, start=2):
            for c, h in enumerate(headers, start=1):
                val = row.get(h, "")
                cell = sheet.cell(r, c, val)
                cell.border = thin
                if h in money_formats and isinstance(val, (int, float)):
                    cell.number_format = '#,##0.00'
        for c, h in enumerate(headers, start=1):
            width = 14
            if h in ("description", "bank_descriptions", "book_descriptions", "reasons"):
                width = 42
            elif h in ("source_file", "match_id"):
                width = 18
            sheet.column_dimensions[get_column_letter(c)].width = width
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows))}"
        sheet.freeze_panes = "A2"

    add_sheet("Transactions", data["transactions"])
    add_sheet("Matches", data["matches"])
    add_sheet("Unmatched Bank", data["unmatched_bank"])
    add_sheet("Unmatched Book", data["unmatched_book"])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
